import os
import io
import re
from typing import Dict, Any, List, Optional
import pymupdf
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# ==============================================================================
# 1. LaTeX & Text Cleaning Utilities
# ==============================================================================

def clean_latex(text: Optional[str]) -> str:
    """Convert LaTeX notation into clean Unicode characters for Word and PDF."""
    if not text:
        return ""
    
    t = str(text)
    # Block & inline math tags
    t = re.sub(r'\$\$(.*?)\$\$', r'\1', t, flags=re.DOTALL)
    t = re.sub(r'\$(.*?)\$', r'\1', t)
    
    # Common fractions
    t = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', t)
    
    # Square roots
    t = re.sub(r'\\sqrt\[([^\]]+)\]\{([^}]+)\}', r'\1√(\2)', t)
    t = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', t)
    
    # Superscripts & powers
    t = re.sub(r'\^2\b', '²', t)
    t = re.sub(r'\^3\b', '³', t)
    t = re.sub(r'\^4\b', '⁴', t)
    t = re.sub(r'\^n\b', 'ⁿ', t)
    t = re.sub(r'\^0\b', '⁰', t)
    t = re.sub(r'\^1\b', '¹', t)
    t = re.sub(r'\^\\circ', '°', t)
    t = re.sub(r'\\degree', '°', t)
    
    # Math Operators & Symbols
    replacements = [
        (r'\\pm', '±'),
        (r'\\mp', '∓'),
        (r'\\times', '×'),
        (r'\\div', '÷'),
        (r'\\neq', '≠'),
        (r'\\leq', '≤'),
        (r'\\geq', '≥'),
        (r'\\approx', '≈'),
        (r'\\in', '∈'),
        (r'\\notin', '∉'),
        (r'\\subset', '⊂'),
        (r'\\subseteq', '⊆'),
        (r'\\cup', '∪'),
        (r'\\cap', '∩'),
        (r'\\emptyset', '∅'),
        (r'\\infty', '∞'),
        (r'\\alpha', 'α'),
        (r'\\beta', 'β'),
        (r'\\gamma', 'γ'),
        (r'\\theta', 'θ'),
        (r'\\pi', 'π'),
        (r'\\omega', 'ω'),
        (r'\\triangle', '△'),
        (r'\\angle', '∠'),
        (r'\\cong', '≅'),
        (r'\\sim', '∼'),
        (r'\\parallel', '∥'),
        (r'\\perp', '⊥'),
        (r'\\Rightarrow', '⇒'),
        (r'\\Leftrightarrow', '⇔'),
        (r'\\to', '→'),
        (r'\\rightarrow', '→'),
        (r'\\quad', '   '),
        (r'\\qquad', '     '),
        (r'\\cdot', '·'),
        (r'\\cdots', '⋯'),
        (r'\\mathbb\{R\}', 'ℝ'),
        (r'\\mathbb\{Z\}', 'ℤ'),
        (r'\\mathbb\{N\}', 'ℕ'),
        (r'\\mathbb\{Q\}', 'ℚ'),
        (r'\\mathbb\{C\}', 'ℂ'),
    ]
    for pattern, repl in replacements:
        t = re.sub(pattern, repl, t)

    # Matrices
    t = re.sub(
        r'\\begin\{bmatrix\}(.*?)\\end\{bmatrix\}',
        lambda m: '[' + m.group(1).replace('\\\\', ' ; ').replace('&', ', ').strip() + ']',
        t,
        flags=re.DOTALL
    )
    t = re.sub(r'\\text\{([^}]+)\}', r'\1', t)
    
    return t.strip()


def has_urdu_text(text: Optional[str]) -> bool:
    """Check if string contains Arabic/Urdu unicode characters."""
    if not text:
        return False
    # Arabic / Urdu unicode range: U+0600 - U+06FF, U+0750 - U+077F, U+FB50 - U+FDFF, U+FE70 - U+FEFF
    return bool(re.search(r'[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]', text))


def normalize_paper_data(paper_data: Dict[str, Any]) -> Dict[str, Any]:
    """Ensure standard keys and types across frontend and database test papers."""
    data = dict(paper_data)
    
    title = data.get("testTitle") or data.get("title") or "Chapter-Wise Test Series 2026"
    campus = data.get("campus") or "Sargodha Road Campus"
    class_name = data.get("className") or data.get("class_name") or "9th"
    subject = data.get("subject") or data.get("subject_name") or "Mathematics"
    syllabus = data.get("syllabus") or data.get("syllabus_name") or "Full Chapter"
    date_val = data.get("date") or "2026-10-17"
    time_allowed = data.get("timeAllowed") or data.get("time_allowed") or "45 Mins"
    instructions = data.get("instructions") or (
        "Note: Attempt all questions according to instructions. "
        "Cutting, overwriting, and pencil usage are prohibited in objective section."
    )
    
    # Questions normalization
    mcqs = []
    short_questions = []
    long_questions = []
    
    # 1. Check sections structure (from GeneratedTestOut)
    sections = data.get("sections")
    if sections and isinstance(sections, list):
        for sec in sections:
            q_list = sec.get("questions") or []
            sec_type = sec.get("question_type") or "short"
            for q in q_list:
                item = {
                    "id": q.get("id"),
                    "question_en": clean_latex(q.get("question_text") or q.get("question_en") or ""),
                    "question_ur": q.get("question_ur") or "",
                    "marks": q.get("marks") or sec.get("marks_per_question") or 1,
                    "options": q.get("options") or [],
                    "correct_answer": q.get("correct_answer"),
                    "solution": q.get("solution") or q.get("explanation") or ""
                }
                if sec_type == "mcq" or q.get("type") == "mcq":
                    mcqs.append(item)
                elif sec_type == "short" or q.get("type") == "short":
                    item["marks"] = item["marks"] or 2
                    short_questions.append(item)
                else:
                    item["marks"] = item["marks"] or 4
                    long_questions.append(item)
    else:
        # 2. Check direct lists from frontend paperData
        raw_mcqs = data.get("mcqs") or []
        for q in raw_mcqs:
            mcqs.append({
                "id": q.get("id"),
                "question_en": clean_latex(q.get("question_en") or q.get("question_text") or ""),
                "question_ur": q.get("question_ur") or "",
                "marks": q.get("marks") or 1,
                "options": q.get("options") or [],
                "correct_option_index": q.get("correct_option_index"),
                "correct_answer": q.get("correct_answer"),
                "solution": q.get("solution") or q.get("explanation") or ""
            })
            
        raw_shorts = data.get("shortQuestions") or data.get("short_questions") or []
        for q in raw_shorts:
            short_questions.append({
                "id": q.get("id"),
                "question_en": clean_latex(q.get("question_en") or q.get("question_text") or ""),
                "question_ur": q.get("question_ur") or "",
                "marks": q.get("marks") or 2,
                "solution": q.get("solution") or q.get("answer") or ""
            })
            
        raw_longs = data.get("longQuestions") or data.get("long_questions") or []
        for q in raw_longs:
            long_questions.append({
                "id": q.get("id"),
                "question_en": clean_latex(q.get("question_en") or q.get("question_text") or ""),
                "question_ur": q.get("question_ur") or "",
                "marks": q.get("marks") or 4,
                "solution": q.get("solution") or q.get("answer") or ""
            })

    # Marks tally
    mcq_total = len(mcqs) * 1
    short_attempt = len(short_questions)
    if len(short_questions) >= 8:
        short_attempt = len(short_questions) - 2
    elif len(short_questions) == 7:
        short_attempt = 6
    elif len(short_questions) == 5:
        short_attempt = 4
    elif len(short_questions) > 3:
        short_attempt = len(short_questions) - 1
        
    short_total = short_attempt * 2
    
    target_marks = data.get("targetMarks") or data.get("total_marks")
    if target_marks == 25 and len(long_questions) >= 2:
        long_total = 8
        long_attempt = 1 if len(long_questions) == 2 else 2
    elif target_marks == 50 and len(long_questions) >= 4:
        long_total = 16
        long_attempt = 2
    elif target_marks == 75 and len(long_questions) >= 6:
        long_total = 24
        long_attempt = 3
    else:
        long_total = sum(q.get("marks", 4) for q in long_questions)
        long_attempt = len(long_questions)
        
    calculated_total = mcq_total + short_total + long_total
    total_marks = target_marks if target_marks else calculated_total

    return {
        "title": title,
        "campus": campus,
        "class_name": class_name,
        "subject": subject,
        "syllabus": syllabus,
        "date": date_val,
        "time_allowed": time_allowed,
        "total_marks": total_marks,
        "instructions": instructions,
        "mcqs": mcqs,
        "short_questions": short_questions,
        "long_questions": long_questions,
        "short_attempt": short_attempt,
        "long_attempt": long_attempt,
        "mcq_total": mcq_total,
        "short_total": short_total,
        "long_total": long_total
    }


# ==============================================================================
# 2. HTML Template Generation (For PyMuPDF Vector PDF)
# ==============================================================================

def generate_test_paper_html(
    paper_data: Dict[str, Any],
    is_answer_key: bool = False,
    language_mode: str = "bilingual"
) -> str:
    """Generate high-fidelity HTML and CSS representation of the test paper or answer key."""
    d = normalize_paper_data(paper_data)
    
    include_en = language_mode in ("bilingual", "en")
    include_ur = language_mode in ("bilingual", "ur")
    
    option_letters = ["A", "B", "C", "D"]
    urdu_option_letters = ["الف", "ب", "ج", "د"]
    
    html = []
    html.append("<!DOCTYPE html>")
    html.append("<html>")
    html.append("<head>")
    html.append("<meta charset='utf-8'>")
    html.append("<style>")
    html.append("""
      @page {
        size: A4 portrait;
        margin: 12mm 12mm 14mm 12mm;
      }
      body {
        font-family: 'Charis SIL', 'Arial', sans-serif;
        color: #0f172a;
        margin: 0;
        padding: 0;
        font-size: 10pt;
        line-height: 1.35;
      }
      .academy-header {
        text-align: center;
        border-bottom: 2.5px solid #0f2744;
        padding-bottom: 6px;
        margin-bottom: 8px;
      }
      .academy-name {
        font-size: 19pt;
        font-weight: 800;
        color: #0f2744;
        text-transform: uppercase;
        letter-spacing: 0.8px;
        margin: 0;
      }
      .academy-sub {
        font-size: 9pt;
        font-weight: 700;
        color: #475569;
        margin-top: 2px;
      }
      .paper-title {
        font-size: 12pt;
        font-weight: 800;
        color: #1e3a8a;
        margin-top: 4px;
        text-transform: uppercase;
      }
      .answer-key-badge {
        display: inline-block;
        background-color: #0f2744;
        color: #ffffff;
        padding: 3px 12px;
        border-radius: 4px;
        font-weight: 800;
        font-size: 10pt;
        margin-top: 4px;
      }
      .meta-table {
        width: 100%;
        border-collapse: collapse;
        margin-bottom: 8px;
        font-size: 9pt;
      }
      .meta-table td {
        border: 1px solid #cbd5e1;
        padding: 4px 6px;
      }
      .meta-label {
        font-weight: 700;
        color: #1e293b;
      }
      .instructions-box {
        background-color: #f8fafc;
        border-left: 3.5px solid #0f2744;
        padding: 4px 8px;
        font-size: 8.5pt;
        color: #334155;
        margin-bottom: 10px;
      }
      .section-banner {
        background-color: #f1f5f9;
        border-top: 1.5px solid #0f2744;
        border-bottom: 1.5px solid #0f2744;
        padding: 4px 8px;
        margin-top: 10px;
        margin-bottom: 8px;
        font-size: 9.5pt;
        font-weight: 800;
        color: #0f2744;
        display: flex;
        justify-content: space-between;
      }
      .q-container {
        margin-bottom: 7px;
        page-break-inside: avoid;
      }
      .q-row {
        display: flex;
        justify-content: space-between;
        align-items: flex-start;
      }
      .q-num {
        font-weight: 800;
        color: #0f2744;
        margin-right: 6px;
        min-width: 24px;
      }
      .q-en {
        flex: 1;
        text-align: left;
        direction: ltr;
        color: #0f172a;
        font-size: 9.5pt;
      }
      .q-ur {
        flex: 1;
        text-align: right;
        direction: rtl;
        font-family: 'Noto Naskh Arabic', 'Jameel Noori Nastaleeq', 'Arial', sans-serif;
        font-size: 10.5pt;
        color: #0f172a;
        margin-top: 1px;
      }
      .q-marks-tag {
        font-weight: 800;
        font-size: 8.5pt;
        color: #475569;
        margin-left: 8px;
        white-space: nowrap;
      }
      .options-table {
        width: 100%;
        border-collapse: collapse;
        margin-top: 3px;
        margin-bottom: 4px;
      }
      .options-table td {
        width: 25%;
        padding: 2px 4px;
        font-size: 8.8pt;
      }
      .opt-tag {
        font-weight: 800;
        color: #1e3a8a;
        margin-right: 3px;
      }
      .key-table {
        width: 100%;
        border-collapse: collapse;
        margin-top: 8px;
        margin-bottom: 12px;
        font-size: 9pt;
      }
      .key-table th {
        background-color: #0f2744;
        color: #ffffff;
        padding: 5px 8px;
        text-align: left;
        font-weight: 700;
      }
      .key-table td {
        border: 1px solid #cbd5e1;
        padding: 5px 8px;
      }
      .key-table tr:nth-child(even) {
        background-color: #f8fafc;
      }
      .correct-pill {
        display: inline-block;
        background-color: #dcfce7;
        color: #15803d;
        border: 1px solid #86efac;
        padding: 1px 6px;
        border-radius: 3px;
        font-weight: 800;
      }
      .solution-box {
        background-color: #f8fafc;
        border-left: 3px solid #3b82f6;
        padding: 4px 8px;
        margin-top: 3px;
        font-size: 8.5pt;
        color: #1e293b;
      }
    """)
    html.append("</style>")
    html.append("</head>")
    html.append("<body>")
    
    # 1. Official Header
    html.append("<div class='academy-header'>")
    html.append("<div class='academy-name'>The Century Academy</div>")
    html.append(f"<div class='academy-sub'>Science • Commerce • Arts | {d['campus']}</div>")
    if is_answer_key:
        html.append("<div class='answer-key-badge'>OFFICIAL ANSWER KEY & MARKING SCHEME</div>")
        html.append(f"<div class='paper-title'>{d['title']}</div>")
    else:
        html.append(f"<div class='paper-title'>{d['title']}</div>")
    html.append("</div>")
    
    # 2. Metadata Grid
    html.append("<table class='meta-table'>")
    html.append("<tr>")
    html.append(f"<td><span class='meta-label'>Class:</span> {d['class_name']}</td>")
    html.append(f"<td><span class='meta-label'>Subject:</span> {d['subject']}</td>")
    html.append(f"<td><span class='meta-label'>Time:</span> {d['time_allowed']}</td>")
    html.append(f"<td><span class='meta-label'>Total Marks:</span> <strong>{d['total_marks']}</strong></td>")
    html.append("</tr>")
    html.append("<tr>")
    html.append(f"<td colspan='2'><span class='meta-label'>Syllabus:</span> {d['syllabus']}</td>")
    html.append(f"<td><span class='meta-label'>Date:</span> {d['date']}</td>")
    if is_answer_key:
        html.append("<td><span class='meta-label'>Type:</span> Solution Guide</td>")
    else:
        html.append("<td><span class='meta-label'>Marks Obt:</span> [ &nbsp;&nbsp;&nbsp;&nbsp; / " + str(d['total_marks']) + " ]</td>")
    html.append("</tr>")
    if not is_answer_key:
        html.append("<tr>")
        html.append("<td colspan='2'><span class='meta-label'>Student Name:</span> ____________________________________</td>")
        html.append("<td colspan='2'><span class='meta-label'>Roll No:</span> _________________ &nbsp;&nbsp; <span class='meta-label'>Sec:</span> ________</td>")
        html.append("</tr>")
    html.append("</table>")
    
    # Instructions Box
    if not is_answer_key and d["instructions"]:
        html.append("<div class='instructions-box'>")
        html.append(f"<strong>Instructions / ہدایات:</strong> {d['instructions']}")
        html.append("</div>")

    # ==========================================================================
    # SECTION A: MCQs
    # ==========================================================================
    if d["mcqs"]:
        html.append("<div class='section-banner'>")
        html.append(f"<span>SECTION A: OBJECTIVE (معروضی حصہ) — Q1</span>")
        html.append(f"<span>({len(d['mcqs'])} × 1 = {d['mcq_total']} Marks)</span>")
        html.append("</div>")
        
        if is_answer_key:
            # Table format for MCQ answer key
            html.append("<table class='key-table'>")
            html.append("<thead><tr><th>Q#</th><th>Question Text</th><th>Correct Option</th><th>Answer Text</th></tr></thead>")
            html.append("<tbody>")
            for idx, q in enumerate(d["mcqs"]):
                # Determine correct option
                c_idx = q.get("correct_option_index")
                if c_idx is None and q.get("correct_answer") is not None:
                    try:
                        c_idx = int(q.get("correct_answer"))
                    except (ValueError, TypeError):
                        pass
                
                c_letter = option_letters[c_idx] if c_idx is not None and 0 <= c_idx < len(option_letters) else "N/A"
                
                # Answer text
                ans_text = ""
                opts = q.get("options") or []
                if c_idx is not None and 0 <= c_idx < len(opts):
                    opt_item = opts[c_idx]
                    ans_text = opt_item.get("text", "") if isinstance(opt_item, dict) else str(opt_item)
                elif q.get("correct_answer"):
                    ans_text = str(q.get("correct_answer"))
                    
                html.append(f"<tr>")
                html.append(f"<td style='font-weight:800; text-align:center;'>{idx+1}</td>")
                html.append(f"<td>{q['question_en']}</td>")
                html.append(f"<td style='text-align:center;'><span class='correct-pill'>Option ({c_letter})</span></td>")
                html.append(f"<td><strong>{clean_latex(ans_text)}</strong></td>")
                html.append(f"</tr>")
            html.append("</tbody></table>")
        else:
            # Paper question format
            for idx, q in enumerate(d["mcqs"]):
                html.append("<div class='q-container'>")
                html.append("<div class='q-row'>")
                html.append(f"<span class='q-num'>({idx+1})</span>")
                if include_en and q["question_en"]:
                    html.append(f"<div class='q-en'>{q['question_en']}</div>")
                html.append("<span class='q-marks-tag'>[1]</span>")
                html.append("</div>")
                
                if include_ur and q["question_ur"]:
                    html.append(f"<div class='q-ur'>{q['question_ur']}</div>")
                
                # Options Grid
                opts = q.get("options") or []
                if opts:
                    html.append("<table class='options-table'><tr>")
                    for o_idx in range(4):
                        opt_val = ""
                        if o_idx < len(opts):
                            opt_val = opts[o_idx].get("text", "") if isinstance(opts[o_idx], dict) else str(opts[o_idx])
                        html.append(f"<td><span class='opt-tag'>({option_letters[o_idx]})</span> {clean_latex(opt_val)}</td>")
                    html.append("</tr></table>")
                html.append("</div>")

    # ==========================================================================
    # SECTION B: Short Questions
    # ==========================================================================
    if d["short_questions"]:
        html.append("<div class='section-banner'>")
        html.append(f"<span>SECTION B: SHORT QUESTIONS (مختصر سوالات) — Q2</span>")
        html.append(f"<span>(Attempt {d['short_attempt']} of {len(d['short_questions'])}) &nbsp; [{d['short_total']} Marks]</span>")
        html.append("</div>")
        
        for idx, q in enumerate(d["short_questions"]):
            html.append("<div class='q-container'>")
            html.append("<div class='q-row'>")
            html.append(f"<span class='q-num'>({idx+1})</span>")
            if include_en and q["question_en"]:
                html.append(f"<div class='q-en'>{q['question_en']}</div>")
            html.append("<span class='q-marks-tag'>[2]</span>")
            html.append("</div>")
            
            if include_ur and q["question_ur"]:
                html.append(f"<div class='q-ur'>{q['question_ur']}</div>")
                
            if is_answer_key:
                sol = q.get("solution") or "Model definition / key steps according to textbook syllabus. 2 Marks full credit for correct formula and steps."
                html.append(f"<div class='solution-box'><strong>Solution / Answer Point:</strong> {clean_latex(sol)}</div>")
            html.append("</div>")

    # ==========================================================================
    # SECTION C: Long Questions
    # ==========================================================================
    if d["long_questions"]:
        html.append("<div class='section-banner'>")
        html.append(f"<span>SECTION C: DETAILED QUESTIONS (انشائیہ حصہ)</span>")
        html.append(f"<span>(Attempt {d['long_attempt']} of {len(d['long_questions'])}) &nbsp; [{d['long_total']} Marks]</span>")
        html.append("</div>")
        
        for idx, q in enumerate(d["long_questions"]):
            q_num_label = f"Q{idx + 3}"
            html.append("<div class='q-container'>")
            html.append("<div class='q-row'>")
            html.append(f"<span class='q-num'><strong>{q_num_label}.</strong></span>")
            if include_en and q["question_en"]:
                html.append(f"<div class='q-en'><strong>{q['question_en']}</strong></div>")
            html.append(f"<span class='q-marks-tag'>[{q.get('marks', 4)}]</span>")
            html.append("</div>")
            
            if include_ur and q["question_ur"]:
                html.append(f"<div class='q-ur'><strong>{q['question_ur']}</strong></div>")
                
            if is_answer_key:
                sol = q.get("solution") or "Detailed step-by-step solution / proof according to Punjab Board / FBISE curriculum. Marks distributed evenly across given, diagram, proof steps, and conclusion."
                html.append(f"<div class='solution-box'><strong>Solution / Marking Criteria:</strong> {clean_latex(sol)}</div>")
            html.append("</div>")

    html.append("</body>")
    html.append("</html>")
    return "\n".join(html)


def generate_test_paper_pdf(
    paper_data: Dict[str, Any],
    is_answer_key: bool = False,
    language_mode: str = "bilingual"
) -> bytes:
    """Generate high-resolution vector PDF using PyMuPDF Story rendering engine."""
    html_content = generate_test_paper_html(paper_data, is_answer_key=is_answer_key, language_mode=language_mode)
    
    story = pymupdf.Story(html=html_content)
    med = pymupdf.paper_rect('a4')
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    
    def rectfn(rect_num, filled):
        # A4 margins: 34pt left/right, 34pt top/bottom
        return med, pymupdf.Rect(34, 34, med.width - 34, med.height - 34), pymupdf.Matrix()
        
    story.write(writer, rectfn)
    writer.close()
    
    return buf.getvalue()


# ==============================================================================
# 3. Microsoft Word (.docx) Document Generator (python-docx)
# ==============================================================================

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in Word table."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex="F1F5F9"):
    """Set cell background color in Word table."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def add_urdu_run(paragraph, text: str, font_size_pt: int = 11, bold: bool = False):
    """Add a bidirectional right-to-left run for Urdu text."""
    if not text:
        return None
    run = paragraph.add_run(text)
    run.font.name = 'Jameel Noori Nastaleeq'
    run.font.size = Pt(font_size_pt)
    run.bold = bold
    
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rtl.set(qn('w:val'), '1')
    rPr.append(rtl)
    return run


def set_paragraph_rtl(paragraph):
    """Mark paragraph as RTL in Word."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)


def generate_test_paper_docx(
    paper_data: Dict[str, Any],
    is_answer_key: bool = False,
    language_mode: str = "bilingual"
) -> bytes:
    """Generate editable Microsoft Word (.docx) document with styling and RTL Urdu support."""
    d = normalize_paper_data(paper_data)
    
    include_en = language_mode in ("bilingual", "en")
    include_ur = language_mode in ("bilingual", "ur")
    option_letters = ["A", "B", "C", "D"]
    
    doc = docx.Document()
    
    # 1. Set Margins (0.5 in / 1.27 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.27)  # A4 width
        section.page_height = Inches(11.69) # A4 height

    # 2. Academy Header
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.paragraph_format.space_after = Pt(2)
    p_head.paragraph_format.space_before = Pt(0)
    
    r_title = p_head.add_run("THE CENTURY ACADEMY")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(15, 39, 68)  # Navy
    r_title.font.name = 'Arial'

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_sub.add_run(f"Science • Commerce • Arts | {d['campus']}")
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    r_sub.font.name = 'Arial'

    p_test = doc.add_paragraph()
    p_test.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_test.paragraph_format.space_after = Pt(6)
    
    if is_answer_key:
        r_badge = p_test.add_run("OFFICIAL ANSWER KEY & MARKING SCHEME\n")
        r_badge.bold = True
        r_badge.font.size = Pt(11)
        r_badge.font.color.rgb = RGBColor(185, 28, 28)
        
    r_test = p_test.add_run(d["title"])
    r_test.bold = True
    r_test.font.size = Pt(12)
    r_test.font.color.rgb = RGBColor(30, 58, 138)
    r_test.font.name = 'Arial'

    # 3. Metadata Table
    meta_table = doc.add_table(rows=3 if not is_answer_key else 2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    widths = [Inches(1.8), Inches(2.2), Inches(1.8), Inches(1.4)]
    for row in meta_table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            set_cell_margins(row.cells[idx], top=60, bottom=60, left=80, right=80)

    # Row 0
    meta_table.cell(0, 0).paragraphs[0].add_run(f"Class: {d['class_name']}").bold = True
    meta_table.cell(0, 1).paragraphs[0].add_run(f"Subject: {d['subject']}").bold = True
    meta_table.cell(0, 2).paragraphs[0].add_run(f"Time: {d['time_allowed']}").bold = True
    meta_table.cell(0, 3).paragraphs[0].add_run(f"Total Marks: {d['total_marks']}").bold = True
    
    # Row 1
    meta_table.cell(1, 0).paragraphs[0].add_run(f"Syllabus: {d['syllabus']}")
    meta_table.cell(1, 1).paragraphs[0].add_run(f"Date: {d['date']}")
    if is_answer_key:
        meta_table.cell(1, 2).paragraphs[0].add_run("Type: Official Key")
        meta_table.cell(1, 3).paragraphs[0].add_run(f"Marks: {d['total_marks']}")
    else:
        meta_table.cell(1, 2).paragraphs[0].add_run("Section: ________")
        meta_table.cell(1, 3).paragraphs[0].add_run(f"Obt Marks: [   /{d['total_marks']}]").bold = True
        
    # Row 2 (Student Name & Roll No)
    if not is_answer_key:
        meta_table.cell(2, 0).paragraphs[0].add_run("Student Name: __________________________________")
        meta_table.cell(2, 2).paragraphs[0].add_run("Roll No: __________________")

    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)
    p_sp.paragraph_format.space_before = Pt(0)
    
    if not is_answer_key and d["instructions"]:
        p_inst = doc.add_paragraph()
        p_inst.paragraph_format.space_after = Pt(6)
        r_inst_lbl = p_inst.add_run("Instructions: ")
        r_inst_lbl.bold = True
        r_inst_lbl.font.size = Pt(8.5)
        r_inst_val = p_inst.add_run(d["instructions"])
        r_inst_val.font.size = Pt(8.5)
        r_inst_val.font.italic = True

    # ==========================================================================
    # SECTION A: MCQs
    # ==========================================================================
    if d["mcqs"]:
        # Banner
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(4)
        r_sec = p_sec.add_run(f"SECTION A: OBJECTIVE (معروضی حصہ) — Q1  ({len(d['mcqs'])} × 1 = {d['mcq_total']} Marks)")
        r_sec.bold = True
        r_sec.font.size = Pt(10)
        r_sec.font.color.rgb = RGBColor(15, 39, 68)
        
        if is_answer_key:
            # Answer key table
            key_tbl = doc.add_table(rows=1, cols=4)
            key_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = key_tbl.rows[0].cells
            hdr[0].paragraphs[0].add_run("Q#").bold = True
            hdr[1].paragraphs[0].add_run("Question Text").bold = True
            hdr[2].paragraphs[0].add_run("Correct Option").bold = True
            hdr[3].paragraphs[0].add_run("Correct Answer").bold = True
            for c in hdr:
                set_cell_shading(c, "0F2744")
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

            for idx, q in enumerate(d["mcqs"]):
                row = key_tbl.add_row().cells
                row[0].paragraphs[0].add_run(str(idx + 1)).bold = True
                row[1].paragraphs[0].add_run(q["question_en"])
                
                c_idx = q.get("correct_option_index")
                if c_idx is None and q.get("correct_answer") is not None:
                    try:
                        c_idx = int(q.get("correct_answer"))
                    except (ValueError, TypeError):
                        pass
                c_letter = option_letters[c_idx] if c_idx is not None and 0 <= c_idx < len(option_letters) else "N/A"
                row[2].paragraphs[0].add_run(f"Option ({c_letter})").bold = True
                
                ans_text = ""
                opts = q.get("options") or []
                if c_idx is not None and 0 <= c_idx < len(opts):
                    opt_item = opts[c_idx]
                    ans_text = opt_item.get("text", "") if isinstance(opt_item, dict) else str(opt_item)
                elif q.get("correct_answer"):
                    ans_text = str(q.get("correct_answer"))
                    
                row[3].paragraphs[0].add_run(clean_latex(ans_text)).bold = True
                for c in row:
                    set_cell_margins(c, top=40, bottom=40, left=60, right=60)
        else:
            # Paper question format
            for idx, q in enumerate(d["mcqs"]):
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(3)
                p_q.paragraph_format.space_after = Pt(2)
                
                r_num = p_q.add_run(f"({idx+1}) ")
                r_num.bold = True
                
                if include_en and q["question_en"]:
                    p_q.add_run(q["question_en"] + " ")
                p_q.add_run("[1]").bold = True
                
                # Urdu question
                if include_ur and q["question_ur"]:
                    p_ur = doc.add_paragraph()
                    p_ur.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    p_ur.paragraph_format.space_before = Pt(0)
                    p_ur.paragraph_format.space_after = Pt(2)
                    set_paragraph_rtl(p_ur)
                    add_urdu_run(p_ur, q["question_ur"], font_size_pt=10.5)

                # Options table
                opts = q.get("options") or []
                if opts:
                    opt_tbl = doc.add_table(rows=1, cols=4)
                    opt_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    row_cells = opt_tbl.rows[0].cells
                    for o_idx in range(4):
                        opt_val = ""
                        if o_idx < len(opts):
                            opt_val = opts[o_idx].get("text", "") if isinstance(opts[o_idx], dict) else str(opts[o_idx])
                        p_opt = row_cells[o_idx].paragraphs[0]
                        p_opt.paragraph_format.space_before = Pt(0)
                        p_opt.paragraph_format.space_after = Pt(0)
                        r_opt_lbl = p_opt.add_run(f"({option_letters[o_idx]}) ")
                        r_opt_lbl.bold = True
                        r_opt_lbl.font.color.rgb = RGBColor(30, 58, 138)
                        
                        clean_opt = clean_latex(opt_val)
                        if has_urdu_text(clean_opt):
                            add_urdu_run(p_opt, clean_opt, font_size_pt=9.5)
                        else:
                            p_opt.add_run(clean_opt)
                        set_cell_margins(row_cells[o_idx], top=20, bottom=20, left=40, right=40)

    # ==========================================================================
    # SECTION B: Short Questions
    # ==========================================================================
    if d["short_questions"]:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(4)
        r_sec = p_sec.add_run(
            f"SECTION B: SHORT QUESTIONS (مختصر سوالات) — Q2  "
            f"(Attempt {d['short_attempt']} of {len(d['short_questions'])}) [{d['short_total']} Marks]"
        )
        r_sec.bold = True
        r_sec.font.size = Pt(10)
        r_sec.font.color.rgb = RGBColor(15, 39, 68)

        for idx, q in enumerate(d["short_questions"]):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(2)
            
            p_q.add_run(f"({idx+1}) ").bold = True
            if include_en and q["question_en"]:
                p_q.add_run(q["question_en"] + " ")
            p_q.add_run("[2]").bold = True
            
            if include_ur and q["question_ur"]:
                p_ur = doc.add_paragraph()
                p_ur.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p_ur.paragraph_format.space_before = Pt(0)
                p_ur.paragraph_format.space_after = Pt(2)
                set_paragraph_rtl(p_ur)
                add_urdu_run(p_ur, q["question_ur"], font_size_pt=10.5)
                
            if is_answer_key:
                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.space_before = Pt(1)
                p_sol.paragraph_format.space_after = Pt(3)
                p_sol.add_run("Solution: ").bold = True
                sol_text = q.get("solution") or "Model definition / key steps according to textbook syllabus. 2 Marks full credit."
                p_sol.add_run(clean_latex(sol_text)).font.italic = True

    # ==========================================================================
    # SECTION C: Long Questions
    # ==========================================================================
    if d["long_questions"]:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(4)
        r_sec = p_sec.add_run(
            f"SECTION C: DETAILED QUESTIONS (انشائیہ حصہ)  "
            f"(Attempt {d['long_attempt']} of {len(d['long_questions'])}) [{d['long_total']} Marks]"
        )
        r_sec.bold = True
        r_sec.font.size = Pt(10)
        r_sec.font.color.rgb = RGBColor(15, 39, 68)

        for idx, q in enumerate(d["long_questions"]):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(2)
            
            p_q.add_run(f"Q{idx + 3}. ").bold = True
            if include_en and q["question_en"]:
                p_q.add_run(q["question_en"] + " ")
            p_q.add_run(f"[{q.get('marks', 4)}]").bold = True
            
            if include_ur and q["question_ur"]:
                p_ur = doc.add_paragraph()
                p_ur.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p_ur.paragraph_format.space_before = Pt(0)
                p_ur.paragraph_format.space_after = Pt(2)
                set_paragraph_rtl(p_ur)
                add_urdu_run(p_ur, q["question_ur"], font_size_pt=10.5, bold=True)
                
            if is_answer_key:
                p_sol = doc.add_paragraph()
                p_sol.paragraph_format.space_before = Pt(1)
                p_sol.paragraph_format.space_after = Pt(3)
                p_sol.add_run("Marking Breakdown & Solution: ").bold = True
                sol_text = q.get("solution") or "Detailed step-by-step solution / proof. Marks distributed across given, diagram, proof steps, and conclusion."
                p_sol.add_run(clean_latex(sol_text)).font.italic = True

    # Save to buffer
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
