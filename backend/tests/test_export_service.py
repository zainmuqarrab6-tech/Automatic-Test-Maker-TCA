import unittest
import io
import docx
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import app
from app.database import Base, get_db
from app.models.academic import Class, Subject, Chapter
from app.models.question import Question
from app.models.scheme import PaperScheme, SchemeSection
from app.models.syllabus import SyllabusDivision, SyllabusDivisionChapter
from app.models.test_paper import GeneratedTest, GeneratedTestQuestion
from app.services.test_exporter import (
    clean_latex,
    has_urdu_text,
    generate_test_paper_pdf,
    generate_test_paper_docx
)

class TestExportService(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine(
            "sqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool
        )
        cls.TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=cls.engine)
        Base.metadata.create_all(bind=cls.engine)

        def override_get_db():
            db = cls.TestingSessionLocal()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = override_get_db
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        Base.metadata.drop_all(bind=cls.engine)
        app.dependency_overrides.clear()

    def setUp(self):
        self.db = self.TestingSessionLocal()
        
        # Populate test academic entities
        c9 = Class(id=1, name="9th")
        self.db.add(c9)
        self.db.flush()

        subj = Subject(id=1, class_id=1, name="Mathematics", code="MATH-9")
        self.db.add(subj)
        self.db.flush()

        ch1 = Chapter(id=1, subject_id=1, chapter_number=1, name="Real Numbers", name_ur="حقیقی اعداد")
        self.db.add(ch1)
        self.db.flush()

        # Questions
        q1 = Question(
            id="q_exp_1", subject_id=1, chapter_id=1, type="mcq",
            question_text="If x^2 = 4, then x is:",
            question_ur="اگر x^2 = 4 ہو تو x کی قیمت ہے:",
            marks=1, difficulty="easy", status="approved",
            options=[{"text": "±2"}, {"text": "2"}, {"text": "-2"}, {"text": "4"}],
            correct_answer="0"
        )
        q2 = Question(
            id="q_exp_2", subject_id=1, chapter_id=1, type="short",
            question_text="Find the conjugate of 3 + 4i.",
            question_ur="3 + 4i کا کاونجوگیٹ معلوم کریں۔",
            marks=2, difficulty="easy", status="approved"
        )
        self.db.add_all([q1, q2])
        self.db.flush()

        # Generated Test record
        test_obj = GeneratedTest(
            id="test-exp-001",
            subject_id=1,
            title="9th Mathematics Midterm Test",
            campus="Main Campus",
            date="2026-10-17",
            time_allowed="45 Mins",
            total_marks=25
        )
        self.db.add(test_obj)
        self.db.flush()

        gtq1 = GeneratedTestQuestion(test_id="test-exp-001", question_id="q_exp_1", order=1)
        gtq2 = GeneratedTestQuestion(test_id="test-exp-001", question_id="q_exp_2", order=2)
        self.db.add_all([gtq1, gtq2])
        self.db.commit()

    def tearDown(self):
        for table in reversed(Base.metadata.sorted_tables):
            self.db.execute(table.delete())
        self.db.commit()
        self.db.close()

    def test_01_clean_latex_and_urdu_detection(self):
        latex_input = r"Solve $\frac{a}{b} \pm \sqrt{x^2} \times \div \neq \leq \degree$"
        cleaned = clean_latex(latex_input)
        self.assertIn("(a / b)", cleaned)
        self.assertIn("±", cleaned)
        self.assertIn("²", cleaned)
        self.assertIn("×", cleaned)
        self.assertIn("÷", cleaned)
        self.assertIn("≠", cleaned)
        self.assertIn("≤", cleaned)
        self.assertIn("°", cleaned)

        # Urdu detection
        self.assertTrue(has_urdu_text("سوال نمبر 1: حل کریں"))
        self.assertFalse(has_urdu_text("Solve Question 1 in English only"))

    def test_02_generate_test_paper_pdf_and_answer_key_pdf(self):
        paper_data = {
            "title": "9th Mathematics Test Paper",
            "campus": "Sargodha Road Campus",
            "className": "9th",
            "subject": "Mathematics",
            "syllabus": "Unit 1: Real Numbers",
            "date": "2026-10-17",
            "timeAllowed": "45 Mins",
            "targetMarks": 25,
            "mcqs": [
                {
                    "id": "q1",
                    "question_en": "Simplify (2 + 3i)(2 - 3i)",
                    "question_ur": "مختصر کریں: (2 + 3i)(2 - 3i)",
                    "options": [{"text": "13"}, {"text": "-5"}, {"text": "5"}, {"text": "1"}],
                    "correct_option_index": 0,
                    "marks": 1
                }
            ],
            "shortQuestions": [
                {
                    "id": "q2",
                    "question_en": "Define additive identity of real numbers.",
                    "question_ur": "حقیقی اعداد کے جمعی ذاتی عنصر کی تعریف کریں۔",
                    "marks": 2,
                    "solution": "Zero (0) is the additive identity because a + 0 = a."
                }
            ],
            "longQuestions": [
                {
                    "id": "q3",
                    "question_en": "Prove that if two lines intersect, vertically opposite angles are equal.",
                    "question_ur": "ثابت کریں کہ جب دو خطوط ایک دوسرے کو قطع کریں تو راسی زاویے برابر ہوتے ہیں۔",
                    "marks": 8,
                    "solution": "Given, To Prove, Construction, and Proof steps with statements and reasons."
                }
            ]
        }

        # 1. Question Paper PDF
        pdf_bytes = generate_test_paper_pdf(paper_data, is_answer_key=False)
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertGreater(len(pdf_bytes), 1000)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

        # 2. Answer Key PDF
        key_pdf_bytes = generate_test_paper_pdf(paper_data, is_answer_key=True)
        self.assertIsInstance(key_pdf_bytes, bytes)
        self.assertGreater(len(key_pdf_bytes), 1000)
        self.assertTrue(key_pdf_bytes.startswith(b"%PDF"))

    def test_03_generate_test_paper_docx_and_answer_key_docx(self):
        paper_data = {
            "title": "Unit 1 Midterm Word Document",
            "campus": "Gulberg Campus",
            "className": "9th",
            "subject": "Mathematics",
            "syllabus": "Unit 1",
            "date": "2026-10-17",
            "timeAllowed": "60 Mins",
            "total_marks": 50,
            "mcqs": [
                {
                    "id": "q1",
                    "question_en": "Matrix with 1 row is called:",
                    "question_ur": "ایک قطار والا قالب کہلاتا ہے:",
                    "options": [{"text": "Row matrix"}, {"text": "Column matrix"}, {"text": "Square"}, {"text": "Null"}],
                    "correct_option_index": 0,
                    "marks": 1
                }
            ],
            "short_questions": [
                {
                    "id": "q2",
                    "question_en": "Find the determinant of matrix A.",
                    "question_ur": "قالب A کا مقطع معلوم کریں۔",
                    "marks": 2,
                    "solution": "|A| = ad - bc = (2)(4) - (1)(3) = 5"
                }
            ]
        }

        # 1. Question Paper DOCX
        docx_bytes = generate_test_paper_docx(paper_data, is_answer_key=False)
        self.assertIsInstance(docx_bytes, bytes)
        self.assertGreater(len(docx_bytes), 5000)
        
        # Verify valid docx by opening with python-docx
        doc = docx.Document(io.BytesIO(docx_bytes))
        self.assertGreaterEqual(len(doc.paragraphs), 3)
        self.assertGreaterEqual(len(doc.tables), 1)

        # 2. Answer Key DOCX
        key_docx_bytes = generate_test_paper_docx(paper_data, is_answer_key=True)
        self.assertIsInstance(key_docx_bytes, bytes)
        self.assertGreater(len(key_docx_bytes), 5000)
        doc_key = docx.Document(io.BytesIO(key_docx_bytes))
        self.assertGreaterEqual(len(doc_key.tables), 2)

    def test_04_export_api_endpoints(self):
        payload = {
            "test_id": "test-exp-001",
            "export_type": "question_paper",
            "language_mode": "bilingual"
        }

        # 1. POST /api/v1/tests/export/pdf
        res = self.client.post("/api/v1/tests/export/pdf", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers.get("content-type"), "application/pdf")
        self.assertTrue(res.content.startswith(b"%PDF"))

        # 2. POST /api/v1/tests/export/docx
        res_docx = self.client.post("/api/v1/tests/export/docx", json=payload)
        self.assertEqual(res_docx.status_code, 200)
        self.assertIn("wordprocessingml", res_docx.headers.get("content-type"))
        self.assertGreater(len(res_docx.content), 5000)

        # 3. POST /api/v1/tests/export/pdf for Answer Key
        payload["export_type"] = "answer_key"
        res_key = self.client.post("/api/v1/tests/export/pdf", json=payload)
        self.assertEqual(res_key.status_code, 200)
        self.assertTrue(res_key.content.startswith(b"%PDF"))

        # 4. GET /api/v1/tests/{test_id}/export/pdf
        res_get_pdf = self.client.get("/api/v1/tests/test-exp-001/export/pdf?export_type=question_paper")
        self.assertEqual(res_get_pdf.status_code, 200)
        self.assertEqual(res_get_pdf.headers.get("content-type"), "application/pdf")

        # 5. GET /api/v1/tests/{test_id}/export/docx
        res_get_docx = self.client.get("/api/v1/tests/test-exp-001/export/docx?export_type=answer_key")
        self.assertEqual(res_get_docx.status_code, 200)
        self.assertIn("wordprocessingml", res_get_docx.headers.get("content-type"))

if __name__ == "__main__":
    unittest.main()
